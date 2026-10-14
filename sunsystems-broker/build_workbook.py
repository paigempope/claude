"""Builds SunSystems_Parts_Desk.xlsx — the parts-broker tracker + quote calculator.

Run:  python build_workbook.py
Then recalc (LibreOffice) so cached values exist for previewers.
"""
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "SunSystems_Parts_Desk.xlsx"
FONT = "Arial"
LOG_ROWS = 300          # pre-built rows in the RFQ log
LIST_ROWS = 150         # pre-built rows in Suppliers / Customers

# ---------- styles ----------
NAVY = "1F3A5F"
INPUT_FILL = PatternFill("solid", fgColor="FFF9C4")   # pale yellow = you type here
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="E8EEF5")
KEY_FILL = PatternFill("solid", fgColor="DFF0D8")
thin = Side(style="thin", color="C9D1DB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_h = Font(name=FONT, size=11, bold=True, color="FFFFFF")
f_sub = Font(name=FONT, size=11, bold=True, color=NAVY)
f_body = Font(name=FONT, size=10)
f_in = Font(name=FONT, size=10, color="0000FF")          # blue = input
f_link = Font(name=FONT, size=10, color="008000")        # green = pulled from another tab
f_note = Font(name=FONT, size=9, italic=True, color="666666")
f_big = Font(name=FONT, size=14, bold=True, color=NAVY)

USD = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "mmm d, yyyy"
STATUSES = ["New", "Sourcing", "Quoted", "Won", "Ordered", "Shipped", "Paid", "Lost"]

S = "Settings"  # sheet-name shortcuts used inside formulas
L = "'RFQ Log'"


def style_range(ws, ref, font=None, fill=None, fmt=None, border=True, align=None):
    for row in ws[ref]:
        for c in row:
            if font:
                c.font = font
            if fill:
                c.fill = fill
            if fmt:
                c.number_format = fmt
            if border:
                c.border = BOX
            if align:
                c.alignment = align


def header(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text)
        c.font = f_h
        c.fill = HEAD_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


wb = Workbook()

# =====================================================================
# 1. START HERE
# =====================================================================
ws = wb.active
ws.title = "Start Here"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 95

ws["B2"] = "SunSystems Parts Desk"
ws["B2"].font = Font(name=FONT, size=20, bold=True, color=NAVY)
ws["B3"] = "Someone needs a part. You know where to get it. This workbook turns that into money, every time."
ws["B3"].font = Font(name=FONT, size=11, italic=True, color="444444")

r = 5
def block(title, rows):
    global r
    ws.cell(row=r, column=2, value=title).font = f_sub
    ws.cell(row=r, column=2).fill = SUB_FILL
    ws.cell(row=r, column=3).fill = SUB_FILL
    r += 1
    for k, v in rows:
        a = ws.cell(row=r, column=2, value=k)
        b = ws.cell(row=r, column=3, value=v)
        a.font = Font(name=FONT, size=10, bold=True)
        b.font = f_body
        b.alignment = Alignment(wrap_text=True, vertical="top")
        a.alignment = Alignment(vertical="top")
        r += 1
    r += 1

block("THE SYSTEM (input → process → output)", [
    ("Input", "A request comes in: part number / spec, quantity, needed-by date. Log it in RFQ Log the same day."),
    ("Process", "Find it, get the supplier's price + shipping, let the sheet compute the quote, send it, collect payment, place the order."),
    ("Output", "Profit. Dashboard shows what you've made, what's in play, and who you owe a follow-up."),
    ("Your edge", "You know where parts live. The customer is paying for your knowledge + time + risk, not for the part. Price accordingly."),
])

block("THE 7 STEPS FOR EVERY REQUEST", [
    ("1. Intake  (New)", "Get EXACT part #, revision, spec, qty, needed-by date, ship-to. A photo of the nameplate/old part beats a description. Log it."),
    ("2. Source  (Sourcing)", "Call 1–3 suppliers. Write down unit cost, shipping, lead time. Pick one. Enter cost in RFQ Log."),
    ("3. Quote  (Quoted)", "Sheet computes Suggested Quote. Override in 'Your Quote' if you know the market. Send a written quote (template in the guide) valid for 7 days."),
    ("4. Close  (Won)", "Customer says yes IN WRITING (email/text is fine). Collect payment or deposit BEFORE you buy."),
    ("5. Order  (Ordered)", "Buy from supplier. Ship direct to the customer when possible (drop-ship) so you never touch inventory."),
    ("6. Deliver (Shipped)", "Send tracking to customer. Confirm it arrived and is the right part."),
    ("7. Get paid (Paid)", "Collect the balance. Enter Date Paid. Profit now shows on the Dashboard."),
])

block("NON-NEGOTIABLE RULES", [
    ("Money first", "Never pay a supplier with your own money before the customer has paid (or paid a deposit that covers your cost). This one rule prevents most losses."),
    ("In writing", "Every quote, every yes, every spec confirmation: text or email. 'Yes, part #XYZ rev C qty 4 at $X' is your contract."),
    ("Exact match", "Wrong part = you eat the return shipping + restocking fee. Confirm part #, revision, voltage, thread, material — whatever matters — before ordering."),
    ("Protect the source", "Don't name your supplier to the customer. Drop-ship with a neutral packing slip when possible. Your contacts are the business."),
    ("Mirror the returns policy", "Your return terms can't be better than your supplier's. If they charge 20% restocking, so do you."),
    ("Min profit floor", "Small orders still cost you an hour. Settings tab enforces a minimum profit per deal (default $50)."),
])

block("HOW TO READ THIS WORKBOOK", [
    ("Yellow cell, blue text", "INPUT — you type here."),
    ("Black text", "FORMULA — don't type over it. It's computed."),
    ("Green text", "Pulled from another tab."),
    ("Rows marked (example)", "Sample data so you can see how it works. Delete or overwrite them when you start."),
])

block("THE TABS", [
    ("Settings", "Your pricing rules: markup tiers by order size, minimum profit, card-fee %, rounding, follow-up days. Change once, every quote follows."),
    ("Quote Calc", "Scratchpad for building one multi-line quote. Gives suggested price, price at a target margin, per-unit customer prices."),
    ("RFQ Log", "The core. One row per request. Track it from New to Paid. Profit, margin, and follow-up dates compute automatically."),
    ("Suppliers", "Your black book. Who carries what, lead times, terms, how reliable they are, how many deals you've done with them."),
    ("Customers", "Who asks, how they found you, what they've bought, what they're worth to you."),
    ("Dashboard", "Profit collected, profit in play, win rate, overdue follow-ups, profit by month."),
])

block("MARKUP vs MARGIN (the engineer's footnote)", [
    ("Markup", "Profit ÷ Cost.   Cost $100, sell $125 → 25% markup."),
    ("Margin", "Profit ÷ Price.  Same deal → $25 ÷ $125 = 20% margin."),
    ("Conversion", "margin = markup ÷ (1 + markup).   markup = margin ÷ (1 − margin).   See the table on the Quote Calc tab."),
    ("Card fees", "If the customer pays by card you lose ~3% of the PRICE. The sheet grosses the quote up so that fee comes out of their price, not your profit."),
])

# =====================================================================
# 2. SETTINGS
# =====================================================================
st = wb.create_sheet("Settings")
st.sheet_view.showGridLines = False
st.column_dimensions["A"].width = 34
st.column_dimensions["B"].width = 14
st.column_dimensions["C"].width = 80
st["A1"] = "Settings — your pricing rules"
st["A1"].font = f_title
st["A2"] = "Change the yellow cells. Every formula in the workbook reads from here."
st["A2"].font = f_note

header(st, 3, ["Setting", "Value", "What it does"])
settings = [
    ("Payment processing fee %", 0.03, PCT,
     "Card/online invoice fee taken from the price (Stripe/Square/QuickBooks ≈ 2.9–3.5%). Set 0 if paid by check/ACH/cash. Assumption — check your processor."),
    ("Minimum profit per deal ($)", 50, USD0,
     "Floor so tiny orders are still worth your time. Assumption — set to what an hour of your time is worth."),
    ("Round quotes UP to nearest ($)", 5, USD0,
     "Clean numbers look professional and add a little cushion. Use 1 for no rounding."),
    ("Quote valid for (days)", 7, "0",
     "Put this on every quote. Supplier prices move; don't get stuck honoring an old one."),
    ("Follow up after (days)", 3, "0",
     "RFQ Log flags the follow-up date this many days after a request comes in."),
    ("Deposit required before ordering %", 0.5, PCT,
     "Minimum up-front payment before you buy. Must at least cover your landed cost. 100% for new customers is reasonable."),
]
for i, (k, v, fmt, note) in enumerate(settings):
    row = 4 + i
    st.cell(row=row, column=1, value=k).font = f_body
    c = st.cell(row=row, column=2, value=v)
    c.font, c.fill, c.number_format = f_in, INPUT_FILL, fmt
    n = st.cell(row=row, column=3, value=note)
    n.font, n.alignment = f_note, Alignment(wrap_text=True, vertical="top")
    for col in (1, 2, 3):
        st.cell(row=row, column=col).border = BOX
    st.row_dimensions[row].height = 28

st["A12"] = "Markup tiers — bigger orders, smaller %"
st["A12"].font = f_sub
header(st, 13, ["Landed cost from ($)", "Markup %", "Why"])
tiers = [
    (0, 0.35, "Small one-off parts: your time is most of the value."),
    (500, 0.25, "Typical repair/replacement order."),
    (2000, 0.18, "Mid-size: customer is more likely to shop around."),
    (10000, 0.12, "Large: still real dollars, keep it competitive."),
    (50000, 0.08, "Very large: negotiate case by case; this is the floor."),
]
for i, (lo, mk, why) in enumerate(tiers):
    row = 14 + i
    a = st.cell(row=row, column=1, value=lo)
    b = st.cell(row=row, column=2, value=mk)
    c = st.cell(row=row, column=3, value=why)
    a.font = b.font = f_in
    a.fill = b.fill = INPUT_FILL
    a.number_format, b.number_format = USD0, PCT
    c.font = f_note
    for col in (1, 2, 3):
        st.cell(row=row, column=col).border = BOX
st["A20"] = ("Tiers are starting-point assumptions, not market data. Keep the 'from' column sorted "
             "smallest → largest, starting at 0. If you know the going rate for a part, override the price in RFQ Log.")
st["A20"].font = f_note
st.merge_cells("A20:C21")
st["A20"].alignment = Alignment(wrap_text=True, vertical="top")

st["A23"] = "Status list (used by dropdowns)"
st["A23"].font = f_sub
for i, s in enumerate(STATUSES):
    c = st.cell(row=24 + i, column=1, value=s)
    c.font, c.border = f_body, BOX
meaning = ["Request logged", "Hunting for supplier/price", "Quote sent, waiting on customer",
           "Customer said yes", "You placed the supplier order", "On its way to customer",
           "Customer paid in full — DONE", "Customer passed / no source"]
for i, m in enumerate(meaning):
    c = st.cell(row=24 + i, column=2, value=m)
    c.font = f_note
    st.merge_cells(start_row=24 + i, start_column=2, end_row=24 + i, end_column=3)

FEE = f"{S}!$B$4"
MINP = f"{S}!$B$5"
ROUND = f"{S}!$B$6"
VALID = f"{S}!$B$7"
FUP = f"{S}!$B$8"
TIER_LO = f"{S}!$A$14:$A$18"
TIER_MK = f"{S}!$B$14:$B$18"

# =====================================================================
# 3. QUOTE CALC
# =====================================================================
q = wb.create_sheet("Quote Calc")
q.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", [18, 34, 8, 14, 14, 16, 16, 4]):
    q.column_dimensions[col].width = w
q["A1"] = "Quote Calculator — build one quote"
q["A1"].font = f_title
q["A2"] = "Enter supplier costs in yellow. The sheet tells you what to charge. Copy the green 'Customer price' column into your quote."
q["A2"].font = f_note

for row, (lbl, val, fmt) in enumerate([
    ("Customer", "Acme Fabrication (example)", None),
    ("Quote #", "Q-001", None),
    ("Date", date.today(), DATE),
], start=4):
    q.cell(row=row, column=1, value=lbl).font = Font(name=FONT, size=10, bold=True)
    c = q.cell(row=row, column=2, value=val)
    c.font, c.fill, c.border = f_in, INPUT_FILL, BOX
    if fmt:
        c.number_format = fmt
q["A7"] = "Valid until"
q["A7"].font = Font(name=FONT, size=10, bold=True)
q["B7"] = f"=B6+{VALID}"
q["B7"].number_format = DATE
q["B7"].font = f_link

header(q, 9, ["Part # / Spec", "Description", "Qty", "Supplier unit cost", "Line cost",
              "Customer unit price", "Customer line total"])
for col, w in zip("ABCDEFG", [18, 34, 8, 14, 14, 16, 16]):
    q.column_dimensions[col].width = w
QL_FIRST, QL_LAST = 10, 19
example_lines = [("6205-2RS", "Sealed ball bearing 25x52x15 (example)", 8, 6.40),
                 ("SC-25-150", "Pneumatic cylinder 25mm bore x 150 stroke (example)", 2, 74.00)]
for row in range(QL_FIRST, QL_LAST + 1):
    for col in range(1, 5):
        c = q.cell(row=row, column=col)
        c.font, c.fill, c.border = f_in, INPUT_FILL, BOX
    q.cell(row=row, column=4).number_format = USD
    if row - QL_FIRST < len(example_lines):
        for col, v in enumerate(example_lines[row - QL_FIRST], start=1):
            q.cell(row=row, column=col, value=v)
    q.cell(row=row, column=5, value=f'=IF(OR(C{row}="",D{row}=""),"",C{row}*D{row})')
    # allocate the final price across lines in proportion to their cost
    q.cell(row=row, column=6,
           value=f'=IF(OR(E{row}="",$E$20=0),"",ROUND(E{row}/$E$20*$E$34/C{row},2))')
    q.cell(row=row, column=7, value=f'=IF(F{row}="","",F{row}*C{row})')
    for col, fmt in ((5, USD), (6, USD), (7, USD)):
        c = q.cell(row=row, column=col)
        c.number_format, c.border = fmt, BOX
        c.font = f_body if col == 5 else Font(name=FONT, size=10, color="008000", bold=True)

def qrow(row, label, formula=None, value=None, fmt=USD, inp=False, bold=False, note=None):
    a = q.cell(row=row, column=4, value=label)
    a.font = Font(name=FONT, size=10, bold=bold)
    a.alignment = Alignment(horizontal="right")
    c = q.cell(row=row, column=5, value=formula if formula is not None else value)
    c.number_format, c.border = fmt, BOX
    if inp:
        c.font, c.fill = f_in, INPUT_FILL
    else:
        c.font = Font(name=FONT, size=10, bold=bold)
    if note:
        n = q.cell(row=row, column=6, value=note)
        n.font = f_note

q.column_dimensions["D"].width = 30
qrow(20, "Parts subtotal", f"=SUM(E{QL_FIRST}:E{QL_LAST})", bold=True)
qrow(21, "Inbound shipping (supplier → you)", value=0, inp=True, note="0 if drop-shipped")
qrow(22, "Outbound shipping (→ customer)", value=18, inp=True, note="example")
qrow(23, "Other costs (fees, cert, packaging)", value=0, inp=True)
qrow(24, "LANDED COST", "=SUM(E20:E23)", bold=True)

qrow(26, "Tier markup (from Settings)", f"=IF(E24=0,0,INDEX({TIER_MK},MATCH(E24,{TIER_LO},1)))", fmt=PCT)
qrow(27, "Suggested quote", f"=IF(E24=0,0,CEILING(MAX(E24*(1+E26),E24+{MINP})/(1-{FEE}),{ROUND}))",
     bold=True, note="tier markup, min-profit floor, card fee covered, rounded up")
qrow(29, "OR: target margin %", value=0.25, fmt=PCT, inp=True, note="what % of the price you want to keep")
qrow(30, "Price at target margin",
     f"=IF(OR(E24=0,1-{FEE}-E29<=0),0,CEILING(E24/(1-{FEE}-E29),{ROUND}))",
     note="price = landed ÷ (1 − fee% − margin%)")
qrow(32, "Which price? (1 = suggested, 2 = target margin, 3 = my own)", value=1, fmt="0", inp=True)
qrow(33, "My own price (if 3)", value=0, inp=True)
qrow(34, "PRICE TO CUSTOMER", "=CHOOSE(E32,E27,E30,E33)", bold=True)
q["E34"].fill = KEY_FILL
q["E34"].font = Font(name=FONT, size=12, bold=True, color=NAVY)
qrow(35, "Card / processing fee", f"=E34*{FEE}")
qrow(36, "YOUR PROFIT", "=E34-E35-E24", bold=True)
q["E36"].fill = KEY_FILL
qrow(37, "Margin (profit ÷ price)", "=IF(E34=0,0,E36/E34)", fmt=PCT)
qrow(38, "Markup (profit ÷ landed cost)", "=IF(E24=0,0,E36/E24)", fmt=PCT)
qrow(39, "Deposit to collect before ordering", f"=MAX(E24,CEILING(E34*{S}!$B$9,1))",
     note="at least your landed cost")
q["F33"] = "only used when 'Which price' = 3"
q["F33"].font = f_note
q["E32"].comment = Comment("1 = Suggested quote (tiers)\n2 = Price at target margin\n3 = My own price (E33)", "Parts Desk")
dv_choice = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=False)
q.add_data_validation(dv_choice)
dv_choice.add("E32")
q.conditional_formatting.add("E36", CellIsRule(operator="lessThan", formula=[MINP],
                             fill=PatternFill("solid", fgColor="F8D7DA")))

# markup ↔ margin cheat table
q["A42"] = "Markup ↔ Margin cheat sheet"
q["A42"].font = f_sub
for i, t in enumerate(["Markup", "→ Margin", "Margin", "→ Markup"], start=1):
    c = q.cell(row=43, column=i, value=t)
    c.font, c.fill, c.border = f_h, HEAD_FILL, BOX
for i, v in enumerate([0.10, 0.15, 0.20, 0.25, 0.30, 0.40, 0.50, 0.75, 1.00]):
    row = 44 + i
    q.cell(row=row, column=1, value=v).number_format = PCT
    q.cell(row=row, column=2, value=f"=A{row}/(1+A{row})").number_format = PCT
    q.cell(row=row, column=3, value=v if v < 1 else 0.5).number_format = PCT
    q.cell(row=row, column=4, value=f"=IF(C{row}>=1,0,C{row}/(1-C{row}))").number_format = PCT
    for col in range(1, 5):
        c = q.cell(row=row, column=col)
        c.border = BOX
        c.font = f_in if col in (1, 3) else f_body
q["C52"] = 0.60

# =====================================================================
# 4. RFQ LOG
# =====================================================================
lg = wb.create_sheet("RFQ Log")
lg.freeze_panes = "E5"
lg["A1"] = "RFQ Log — every request, New → Paid"
lg["A1"].font = f_title
lg["A2"] = ("Type in yellow columns only. Leave 'Your Quote' blank to use the Suggested Quote. "
            "Fill 'Date Paid' when the money lands — that's what the Dashboard counts.")
lg["A2"].font = f_note
cols = [
    ("ID", 8), ("Date In", 12), ("Customer", 22), ("Part # / Spec", 18), ("Description", 30),
    ("Qty", 6), ("Needed By", 12), ("Supplier", 20), ("Unit Cost ($)", 11),
    ("Shipping + Fees ($)", 11), ("Landed Cost ($)", 12), ("Tier Markup", 9),
    ("Suggested Quote ($)", 12), ("Your Quote ($)", 12), ("Price Sent ($)", 12),
    ("Est. Profit ($)", 12), ("Margin", 8), ("Status", 11), ("Follow Up On", 12),
    ("Days Open", 8), ("Date Paid", 12), ("Notes", 40),
]
header(lg, 4, [c[0] for c in cols], [c[1] for c in cols])
INPUT_COLS = {2, 3, 4, 5, 6, 7, 8, 9, 10, 14, 18, 21, 22}
FIRST = 5
LAST = FIRST + LOG_ROWS - 1
today = date.today()
examples = [
    (today - timedelta(days=40), "Acme Fabrication (example)", "6205-2RS", "Sealed ball bearing (example)", 8,
     today - timedelta(days=30), "Midwest Bearing Supply (example)", 6.40, 18, None, "Paid",
     today - timedelta(days=28), "Repeat buyer — ask about monthly stock"),
    (today - timedelta(days=12), "Lakeside Machine Shop (example)", "VFD-7.5HP-480V", "Variable frequency drive (example)", 1,
     today + timedelta(days=5), "Great Lakes Industrial (example)", 1180, 45, 1650, "Shipped",
     None, "Customer quoted $1,900 elsewhere; priced under it"),
    (today - timedelta(days=4), "Acme Fabrication (example)", "SC-25-150", "Pneumatic cylinder (example)", 2,
     today + timedelta(days=14), "Midwest Bearing Supply (example)", 74, 15, None, "Quoted",
     None, "Waiting on PO"),
]
for row in range(FIRST, LAST + 1):
    r_ = str(row)
    f = {
        1: f'=IF(B{r_}="","","Q-"&TEXT(ROW()-{FIRST - 1},"000"))',
        11: f'=IF(OR(F{r_}="",I{r_}=""),"",F{r_}*I{r_}+N(J{r_}))',
        12: f'=IF(K{r_}="","",INDEX({TIER_MK},MATCH(K{r_},{TIER_LO},1)))',
        13: f'=IF(K{r_}="","",CEILING(MAX(K{r_}*(1+L{r_}),K{r_}+{MINP})/(1-{FEE}),{ROUND}))',
        15: f'=IF(K{r_}="","",IF(N{r_}="",M{r_},N{r_}))',
        16: f'=IF(O{r_}="","",O{r_}*(1-{FEE})-K{r_})',
        17: f'=IF(OR(O{r_}="",O{r_}=0),"",P{r_}/O{r_})',
        19: f'=IF(OR(B{r_}="",R{r_}="Paid",R{r_}="Lost",R{r_}="Shipped"),"",B{r_}+{FUP})',
        20: f'=IF(OR(B{r_}="",R{r_}="Paid",R{r_}="Lost"),"",TODAY()-B{r_})',
    }
    for col in range(1, len(cols) + 1):
        c = lg.cell(row=row, column=col)
        c.border = BOX
        if col in INPUT_COLS:
            c.font, c.fill = f_in, INPUT_FILL
        else:
            c.font = f_body
        if col in f:
            c.value = f[col]
        if col in (2, 7, 19, 21):
            c.number_format = DATE
        elif col in (9, 10, 11, 13, 14, 15, 16):
            c.number_format = USD
        elif col in (12, 17):
            c.number_format = PCT
    lg.cell(row=row, column=22).alignment = Alignment(wrap_text=False)
for i, ex in enumerate(examples):
    row = FIRST + i
    (d_in, cust, part, desc, qty, need, sup, cost, ship, override, status, paid, note) = ex
    for col, v in ((2, d_in), (3, cust), (4, part), (5, desc), (6, qty), (7, need), (8, sup),
                   (9, cost), (10, ship), (14, override), (18, status), (21, paid), (22, note)):
        lg.cell(row=row, column=col, value=v)

dv_status = DataValidation(type="list", formula1=f"={S}!$A$24:$A$31", allow_blank=True)
dv_cust = DataValidation(type="list", formula1=f"=Customers!$A$5:$A${4 + LIST_ROWS}", allow_blank=True,
                         showErrorMessage=False)
dv_sup = DataValidation(type="list", formula1=f"=Suppliers!$A$5:$A${4 + LIST_ROWS}", allow_blank=True,
                        showErrorMessage=False)
for dv, col in ((dv_status, "R"), (dv_cust, "C"), (dv_sup, "H")):
    lg.add_data_validation(dv)
    dv.add(f"{col}{FIRST}:{col}{LAST}")

status_colors = {"New": "E3F2FD", "Sourcing": "FFF3E0", "Quoted": "FFF9C4", "Won": "C8E6C9",
                 "Ordered": "B2DFDB", "Shipped": "B3E5FC", "Paid": "A5D6A7", "Lost": "E0E0E0"}
rng = f"R{FIRST}:R{LAST}"
for s_, color in status_colors.items():
    lg.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{s_}"'],
                                  fill=PatternFill("solid", fgColor=color)))
# thin margin warning, overdue follow-ups
lg.conditional_formatting.add(f"Q{FIRST}:Q{LAST}",
    FormulaRule(formula=[f'AND(Q{FIRST}<>"",Q{FIRST}<0.1)'], fill=PatternFill("solid", fgColor="F8D7DA")))
lg.conditional_formatting.add(f"S{FIRST}:S{LAST}",
    FormulaRule(formula=[f'AND(S{FIRST}<>"",S{FIRST}<=TODAY(),OR(R{FIRST}="New",R{FIRST}="Sourcing",R{FIRST}="Quoted"))'],
                fill=PatternFill("solid", fgColor="FFCC80"), font=Font(name=FONT, bold=True, color="8A4B00")))
lg["N4"].comment = Comment("Leave blank to send the Suggested Quote. Type a number to override "
                           "(e.g. you know the customer's other quote).", "Parts Desk")
lg["J4"].comment = Comment("Everything that's a cost to you besides the part itself: inbound + outbound "
                           "shipping, hazmat, cert fees, packaging.", "Parts Desk")
lg["P4"].comment = Comment("Price × (1 − card fee %) − landed cost.", "Parts Desk")
lg.auto_filter.ref = f"A4:V{LAST}"

# =====================================================================
# 5. SUPPLIERS
# =====================================================================
sp = wb.create_sheet("Suppliers")
sp.freeze_panes = "B5"
sp["A1"] = "Suppliers — the black book"
sp["A1"].font = f_title
sp["A2"] = "Every source you know. 'Deals' and 'Profit' fill in automatically from RFQ Log."
sp["A2"].font = f_note
scols = [("Supplier", 26), ("Contact person", 18), ("Phone", 15), ("Email", 26),
         ("What they carry", 32), ("Lead time (days)", 10), ("Payment terms", 14),
         ("Min order / notes", 18), ("Drop-ship?", 9), ("Rating (1–5)", 8),
         ("Deals", 8), ("Profit through them ($)", 14), ("Notes", 36)]
header(sp, 4, [c[0] for c in scols], [c[1] for c in scols])
sup_examples = [
    ("Midwest Bearing Supply (example)", "Dana", "555-0100", "sales@example.com",
     "Bearings, seals, belts, PT components", 2, "Net 30 after 1st order", "$50 min", "Yes", 5,
     "Fast, ask for counter price"),
    ("Great Lakes Industrial (example)", "Sam", "555-0144", "quotes@example.com",
     "Drives, motors, controls", 5, "Prepay", "None", "Yes", 4, "Call, don't email"),
]
for row in range(5, 5 + LIST_ROWS):
    for col in range(1, len(scols) + 1):
        c = sp.cell(row=row, column=col)
        c.border = BOX
        if col in (11, 12):
            c.font = f_link
        else:
            c.font, c.fill = f_in, INPUT_FILL
    sp.cell(row=row, column=11, value=f'=IF(A{row}="","",COUNTIF({L}!$H${FIRST}:$H${LAST},A{row}))')
    sp.cell(row=row, column=12,
            value=f'=IF(A{row}="","",SUMIFS({L}!$P${FIRST}:$P${LAST},{L}!$H${FIRST}:$H${LAST},A{row},{L}!$R${FIRST}:$R${LAST},"Paid"))')
    sp.cell(row=row, column=12).number_format = USD
for i, ex in enumerate(sup_examples):
    for col, v in zip([1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 13], ex):
        sp.cell(row=5 + i, column=col, value=v)
dv_yn = DataValidation(type="list", formula1='"Yes,No,Ask"', allow_blank=True)
dv_rate = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
sp.add_data_validation(dv_yn)
sp.add_data_validation(dv_rate)
dv_yn.add(f"I5:I{4 + LIST_ROWS}")
dv_rate.add(f"J5:J{4 + LIST_ROWS}")

# =====================================================================
# 6. CUSTOMERS
# =====================================================================
cu = wb.create_sheet("Customers")
cu.freeze_panes = "B5"
cu["A1"] = "Customers — who needs parts"
cu["A1"].font = f_title
cu["A2"] = "Name must match exactly what you pick in RFQ Log. Totals fill in automatically."
cu["A2"].font = f_note
ccols = [("Customer / Company", 28), ("Contact person", 18), ("Phone", 15), ("Email", 26),
         ("How they found you", 18), ("Payment terms you give", 16), ("Requests", 9),
         ("Won", 7), ("Win rate", 8), ("Revenue paid ($)", 13), ("Profit paid ($)", 13), ("Notes", 36)]
header(cu, 4, [c[0] for c in ccols], [c[1] for c in ccols])
cust_examples = [
    ("Acme Fabrication (example)", "Jordan", "555-0199", "buyer@example.com", "Referral", "50% deposit",
     "Maintenance mgr; orders bearings monthly"),
    ("Lakeside Machine Shop (example)", "Chris", "555-0123", "office@example.com", "Word of mouth", "Prepay",
     "Downtime-driven — pays for speed"),
]
WON = '{"Won","Ordered","Shipped","Paid"}'
for row in range(5, 5 + LIST_ROWS):
    for col in range(1, len(ccols) + 1):
        c = cu.cell(row=row, column=col)
        c.border = BOX
        if col in (7, 8, 9, 10, 11):
            c.font = f_link
        else:
            c.font, c.fill = f_in, INPUT_FILL
    cu.cell(row=row, column=7, value=f'=IF(A{row}="","",COUNTIF({L}!$C${FIRST}:$C${LAST},A{row}))')
    cu.cell(row=row, column=8,
            value=f'=IF(A{row}="","",' + "+".join(f'COUNTIFS({L}!$C${FIRST}:$C${LAST},A{row},{L}!$R${FIRST}:$R${LAST},"{st}")' for st in ("Won","Ordered","Shipped","Paid")) + ')')
    cu.cell(row=row, column=9, value=f'=IF(OR(A{row}="",G{row}=0),"",H{row}/G{row})')
    cu.cell(row=row, column=10,
            value=f'=IF(A{row}="","",SUMIFS({L}!$O${FIRST}:$O${LAST},{L}!$C${FIRST}:$C${LAST},A{row},{L}!$R${FIRST}:$R${LAST},"Paid"))')
    cu.cell(row=row, column=11,
            value=f'=IF(A{row}="","",SUMIFS({L}!$P${FIRST}:$P${LAST},{L}!$C${FIRST}:$C${LAST},A{row},{L}!$R${FIRST}:$R${LAST},"Paid"))')
    cu.cell(row=row, column=9).number_format = PCT
    cu.cell(row=row, column=10).number_format = USD
    cu.cell(row=row, column=11).number_format = USD
for i, ex in enumerate(cust_examples):
    for col, v in zip([1, 2, 3, 4, 5, 6, 12], ex):
        cu.cell(row=5 + i, column=col, value=v)
dv_src = DataValidation(type="list", formula1='"Referral,Word of mouth,Repeat,Online,Cold call,Trade show,Other"',
                        allow_blank=True, showErrorMessage=False)
cu.add_data_validation(dv_src)
dv_src.add(f"E5:E{4 + LIST_ROWS}")

# =====================================================================
# 7. DASHBOARD
# =====================================================================
db = wb.create_sheet("Dashboard")
db.sheet_view.showGridLines = False
for col, w in zip("ABCDEFG", [34, 16, 4, 24, 14, 14, 4]):
    db.column_dimensions[col].width = w
db["A1"] = "Dashboard"
db["A1"].font = f_title
db["A2"] = "=\"As of \"&TEXT(TODAY(),\"mmm d, yyyy\")"
db["A2"].font = f_note

R_ = f"{L}!$R${FIRST}:$R${LAST}"
O_ = f"{L}!$O${FIRST}:$O${LAST}"
P_ = f"{L}!$P${FIRST}:$P${LAST}"
B_ = f"{L}!$B${FIRST}:$B${LAST}"
S_ = f"{L}!$S${FIRST}:$S${LAST}"
U_ = f"{L}!$U${FIRST}:$U${LAST}"
WONC = "+".join(f'COUNTIF({R_},"{st}")' for st in ("Won","Ordered","Shipped","Paid"))

kpis = [
    ("MONEY", None, None),
    ("Profit collected (Paid)", f'=SUMIFS({P_},{R_},"Paid")', USD),
    ("Revenue collected (Paid)", f'=SUMIFS({O_},{R_},"Paid")', USD),
    ("Average margin on paid deals", "=IF(B5=0,0,B4/B5)", PCT),
    ("Profit in play (Won/Ordered/Shipped)", f'=SUMIFS({P_},{R_},"Won")+SUMIFS({P_},{R_},"Ordered")+SUMIFS({P_},{R_},"Shipped")', USD),
    ("Profit on the table (Quoted)", f'=SUMIFS({P_},{R_},"Quoted")', USD),
    ("Profit this month (by Date Paid)",
     f'=SUMIFS({P_},{R_},"Paid",{U_},">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),{U_},"<"&EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),1))', USD),
    ("", None, None),
    ("ACTIVITY", None, None),
    ("Total requests logged", f'=COUNT({B_})', "0"),
    ("Win rate (won ÷ decided)",
     f'=IFERROR(({WONC})/(({WONC})+COUNTIF({R_},"Lost")),0)', PCT),
    ("Open requests (New/Sourcing/Quoted)", f'=COUNTIF({R_},"New")+COUNTIF({R_},"Sourcing")+COUNTIF({R_},"Quoted")', "0"),
    ("FOLLOW-UPS DUE NOW",
     f'=COUNTIFS({S_},"<="&TODAY(),{R_},"New")+COUNTIFS({S_},"<="&TODAY(),{R_},"Sourcing")+COUNTIFS({S_},"<="&TODAY(),{R_},"Quoted")', "0"),
    ("Shipped, waiting on payment", f'=COUNTIF({R_},"Shipped")', "0"),
]
for i, (lbl, formula, fmt) in enumerate(kpis):
    row = 3 + i
    a = db.cell(row=row, column=1, value=lbl)
    if formula is None:
        a.font = f_sub
        if lbl:
            a.fill = SUB_FILL
            db.cell(row=row, column=2).fill = SUB_FILL
        continue
    a.font = f_body
    a.border = BOX
    c = db.cell(row=row, column=2, value=formula)
    c.number_format, c.font, c.border = fmt, f_big, BOX
db.cell(row=4, column=2).fill = KEY_FILL
db.conditional_formatting.add("B15", CellIsRule(operator="greaterThan", formula=["0"],
                              fill=PatternFill("solid", fgColor="FFCC80")))

db["D3"] = "PIPELINE BY STATUS"
db["D3"].font = f_sub
for col in "DEF":
    db[f"{col}3"].fill = SUB_FILL
header_cells = [("D4", "Status"), ("E4", "Count"), ("F4", "Price sent ($)")]
for ref, t in header_cells:
    db[ref] = t
    db[ref].font, db[ref].fill, db[ref].border = f_h, HEAD_FILL, BOX
for i, s_ in enumerate(STATUSES):
    row = 5 + i
    db.cell(row=row, column=4, value=s_).font = f_body
    db.cell(row=row, column=5, value=f'=COUNTIF({R_},D{row})').number_format = "0"
    c = db.cell(row=row, column=6, value=f'=SUMIFS({O_},{R_},D{row})')
    c.number_format = USD0
    for col in (4, 5, 6):
        db.cell(row=row, column=col).border = BOX
        db.cell(row=row, column=col).fill = PatternFill("solid", fgColor=status_colors[s_])

db["A19"] = "PROFIT BY MONTH (this year, by Date Paid)"
db["A19"].font = f_sub
db["A19"].fill = SUB_FILL
db["B19"].fill = SUB_FILL
for ref, t in (("A20", "Month"), ("B20", "Profit ($)")):
    db[ref] = t
    db[ref].font, db[ref].fill, db[ref].border = f_h, HEAD_FILL, BOX
for m in range(1, 13):
    row = 20 + m
    a = db.cell(row=row, column=1, value=f"=DATE(YEAR(TODAY()),{m},1)")
    a.number_format = "mmm yyyy"
    b = db.cell(row=row, column=2,
                value=f'=SUMIFS({P_},{R_},"Paid",{U_},">="&A{row},{U_},"<"&EDATE(A{row},1))')
    b.number_format = USD0
    for c in (a, b):
        c.border = BOX
        c.font = f_body
chart = BarChart()
chart.type = "col"
chart.title = "Profit by month"
chart.y_axis.title = "Profit ($)"
chart.y_axis.numFmt = "$#,##0"
chart.legend = None
chart.add_data(Reference(db, min_col=2, min_row=20, max_row=32), titles_from_data=True)
chart.set_categories(Reference(db, min_col=1, min_row=21, max_row=32))
chart.x_axis.number_format = "mmm"
chart.height, chart.width = 8, 16
chart.series[0].graphicalProperties.solidFill = NAVY
db.add_chart(chart, "D19")

# workbook-wide: default font for any untouched cell
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                c.font = Font(name=FONT, size=c.font.size if c.font else 10, bold=c.font.bold if c.font else False,
                              italic=c.font.italic if c.font else False, color=c.font.color if c.font else None)

wb.move_sheet("Dashboard", offset=-5)  # Start Here, Dashboard, Settings, ...
wb.active = 0
wb.save(OUT)
print("wrote", OUT)
